"""M5 Excel check: treaty mechanics in live formulas, recalculated by LibreOffice headless and compared with Python.

Sheet `toy`: the SPEC §11 toy through 50 xs 60 (c = 1, n = 1).
Sheet `history`: calibration as-if events through two sample layers (q = 0, c = 1, n = 1):
attach 1-in-10 to exhaust 1-in-50, and attach 1-in-50 to exhaust 1-in-200 (dollar terms from the base OEP, rounded).
"""
from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from floodcat.reinsurance.engine import Layer, Programme, YearEventTable, apply_programme
from floodcat.utils import risk_measures as rm
from floodcat.utils.io import EXCEL_DIR, PROCESSED_DIR, TABLES_DIR, load_config, write_json

BOLD = Font(bold=True)
INPUT = PatternFill("solid", fgColor="FFF2CC")
TOY = [(1, 12), (3, 65), (3, 8), (4, 30), (6, 140), (7, 20), (8, 55), (8, 90), (10, 5)]


def _write_layer_block(ws, col0: int, title: str, attach: float, limit: float, c: float, n: int,
                       ev_rows: tuple[int, int], ev_year_col: str, ev_loss_col: str, years: list[int], top: int, rec_col: int) -> dict:
    """Parameters, per-event recovery column and per-year table for one layer. Returns cell addresses."""
    L = get_column_letter
    ws.cell(top, col0, title).font = BOLD
    names = ["Attachment A", "Limit L", "Placement c", "Reinstatements n", "Reinstatement rate"]
    vals = [attach, limit, c, n, 1.0]
    for i, (k, v) in enumerate(zip(names, vals)):
        ws.cell(top + 1 + i, col0, k)
        cell = ws.cell(top + 1 + i, col0 + 1, v)
        cell.fill = INPUT
    a, lim, cc, nn = (f"${L(col0 + 1)}${top + 1 + i}" for i in range(4))
    # per-event recovery next to the event list
    r0, r1 = ev_rows
    ws.cell(r0 - 1, rec_col, f"r_e ({title})").font = BOLD
    for r in range(r0, r1 + 1):
        ws.cell(r, rec_col, f"={cc}*MIN(MAX({ev_loss_col}{r}-{a},0),{lim})")
    # per-year table
    y0 = top + 8
    hdr = ["Year", "Sum r_e", "Annual cap (1+n)cL", "Recovery R", "Reinstated f"]
    for j, h in enumerate(hdr):
        ws.cell(y0, col0 + j, h).font = BOLD
    rc = L(rec_col)
    for i, y in enumerate(years, 1):
        row = y0 + i
        ws.cell(row, col0, y)
        ws.cell(row, col0 + 1, f"=SUMIF(${ev_year_col}${r0}:${ev_year_col}${r1},{L(col0)}{row},${rc}${r0}:${rc}${r1})")
        ws.cell(row, col0 + 2, f"=(1+{nn})*{cc}*{lim}")
        ws.cell(row, col0 + 3, f"=MIN({L(col0 + 1)}{row},{L(col0 + 2)}{row})")
        ws.cell(row, col0 + 4, f"=MIN({L(col0 + 3)}{row}/({cc}*{lim}),{nn})")
    last = y0 + len(years)
    ws.cell(last + 1, col0, "Burning cost").font = BOLD
    ws.cell(last + 1, col0 + 3, f"=AVERAGE({L(col0 + 3)}{y0 + 1}:{L(col0 + 3)}{last})")
    ws.cell(last + 2, col0, "Years hit").font = BOLD
    ws.cell(last + 2, col0 + 3, f'=COUNTIF({L(col0 + 3)}{y0 + 1}:{L(col0 + 3)}{last},">0")')
    ws.cell(last + 3, col0, "Mean f").font = BOLD
    ws.cell(last + 3, col0 + 4, f"=AVERAGE({L(col0 + 4)}{y0 + 1}:{L(col0 + 4)}{last})")
    return {"R": (L(col0 + 3), y0 + 1, last), "f": (L(col0 + 4), y0 + 1, last), "bc": f"{L(col0 + 3)}{last + 1}",
            "hit": f"{L(col0 + 3)}{last + 2}"}


def sample_layers() -> list[tuple[str, float, float]]:
    ri = load_config("reinsurance")
    ylt = pd.read_parquet(PROCESSED_DIR / "ylt.parquet")
    rnd = ri["rounding"]
    lv = lambda rp: float(rnd * round(rm.return_level(ylt["max_event"].to_numpy(), rp) / rnd))
    return [("Layer 1-in-10 to 1-in-50", lv(10), lv(50) - lv(10)), ("Layer 1-in-50 to 1-in-200", lv(50), lv(200) - lv(50))]


def build(path: Path) -> dict:
    """Write the workbook; return the Python values it must reproduce, keyed by sheet!cell."""
    from floodcat.capital.run import history_table

    wb = openpyxl.Workbook()
    expected: dict = {}

    # --- toy
    ws = wb.active
    ws.title = "toy"
    ws["A1"] = "SPEC §11 toy YELT (USD m, attritional 0) through 50 xs 60"
    ws["A1"].font = BOLD
    ws.cell(3, 1, "Year").font = BOLD
    ws.cell(3, 2, "Event loss S").font = BOLD
    for i, (y, s) in enumerate(TOY, 4):
        ws.cell(i, 1, y)
        ws.cell(i, 2, s)
    blk = _write_layer_block(ws, 5, "50 xs 60", 60.0, 50.0, 1.0, 1, (4, 3 + len(TOY)), "A", "B", list(range(1, 11)), 2, rec_col=3)
    t = YearEventTable(np.array([s for _, s in TOY], float), np.array([y - 1 for y, _ in TOY]), np.zeros(10))
    res = apply_programme(t, Programme("toy", 0.0, (Layer("L", 60.0, 50.0, 1.0, 1),)))["layers"]["L"]
    _expect(expected, "toy", blk, res)

    # --- history
    port = load_config("portfolio")
    years = list(range(port["calibration"]["start"], port["calibration"]["end"] + 1))
    ev = pd.read_parquet(PROCESSED_DIR / "event_catalogue.parquet")
    inc = ev[ev["included"] & ev["year"].isin(years)].sort_values(["year", "first_date"])
    ws = wb.create_sheet("history")
    ws["A1"] = "Calibration as-if events (USD m, 2025 basis) through two sample layers, q = 0, c = 1, n = 1"
    ws["A1"].font = BOLD
    for j, h in enumerate(["Year", "Event", "As-if loss S"], 1):
        ws.cell(3, j, h).font = BOLD
    for i, r in enumerate(inc.itertuples(), 4):
        ws.cell(i, 1, int(r.year))
        ws.cell(i, 2, r.name)
        ws.cell(i, 3, round(float(r.loss_asif), 6))
    r0, r1 = 4, 3 + len(inc)
    ht = history_table()
    ht = YearEventTable(np.round(ht.event_loss, 6), ht.year_id, ht.gross)
    col0 = 7
    for k, (title, a, lim) in enumerate(sample_layers()):
        blk = _write_layer_block(ws, col0, title, a, lim, 1.0, 1, (r0, r1), "A", "C", years, 2, rec_col=4 + k)
        res = apply_programme(ht, Programme("h", 0.0, (Layer("L", a, lim, 1.0, 1),)))["layers"]["L"]
        _expect(expected, "history", blk, res)
        col0 += 7
    for w in wb.worksheets:
        w.column_dimensions["B"].width = 28
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return expected


def _expect(expected: dict, sheet: str, blk: dict, res: dict) -> None:
    col, a, b = blk["R"]
    for i, v in enumerate(res["recovery"]):
        expected[f"{sheet}!{col}{a + i}"] = float(v)
    col, a, b = blk["f"]
    for i, v in enumerate(res["f"]):
        expected[f"{sheet}!{col}{a + i}"] = float(v)
    expected[f"{sheet}!{blk['bc']}"] = float(res["recovery"].mean())
    expected[f"{sheet}!{blk['hit']}"] = float((res["recovery"] > 0).sum())


def recalc(path: Path) -> Path:
    """Recalculate with LibreOffice headless; returns the recalculated copy."""
    soffice = shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    out = Path(tempfile.mkdtemp())
    subprocess.run([soffice, "--headless", "--calc", "--convert-to", "xlsx", "--outdir", str(out), str(path)],
                   check=True, capture_output=True, timeout=180)
    return out / path.name


def compare(recalculated: Path, expected: dict) -> dict:
    wb = openpyxl.load_workbook(recalculated, data_only=True)
    worst, rows = 0.0, []
    for key, py in expected.items():
        sheet, cell = key.split("!")
        xl = wb[sheet][cell].value
        rel = abs(xl - py) / max(abs(py), 1e-9) if py != 0 else abs(xl)
        worst = max(worst, rel)
        rows.append({"cell": key, "python": py, "excel": xl, "rel_diff": rel})
    return {"n_cells": len(rows), "max_rel_diff": worst, "tolerance": 1e-4, "pass": bool(worst <= 1e-4), "cells": rows}


def run() -> None:
    path = EXCEL_DIR / "reinsurance_checks.xlsx"
    expected = build(path)
    result = compare(recalc(path), expected)
    write_json(result, TABLES_DIR / "excel_check.json")
    print(f"Excel check: {result['n_cells']} cells, max rel diff {result['max_rel_diff']:.2e}, {'PASS' if result['pass'] else 'FAIL'}", flush=True)
